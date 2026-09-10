"""Regenerate the small, synthetic Milestone 2 parser fixtures."""

from __future__ import annotations

import shutil
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_COLOR_INDEX
from pypdf import PdfWriter

FIXTURE_ROOT = Path(__file__).parent / "library"
STRUCTURE_FIXTURE_ROOT = Path(__file__).parent / "structure"


def _save_document(
    relative_path: str,
    build: Callable[[DocumentObject], None],
) -> None:
    path = FIXTURE_ROOT / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    build(document)
    document.save(path)


def _save_structure_document(
    filename: str,
    build: Callable[[DocumentObject], None],
) -> None:
    path = STRUCTURE_FIXTURE_ROOT / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    build(document)
    document.save(path)


def _write_text_pdf(path: Path, text: str) -> None:
    """Write a tiny one-page PDF with an extractable Helvetica text stream."""
    path.parent.mkdir(parents=True, exist_ok=True)
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]

    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{object_number} 0 obj\n".encode())
        content.extend(body)
        content.extend(b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode())
    content.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    path.write_bytes(content)


def generate() -> None:
    if FIXTURE_ROOT.exists():
        shutil.rmtree(FIXTURE_ROOT)
    if STRUCTURE_FIXTURE_ROOT.exists():
        shutil.rmtree(STRUCTURE_FIXTURE_ROOT)

    def ordinary(document: DocumentObject) -> None:
        document.add_paragraph("Governments should invest in resilient public transit.")
        document.add_paragraph("Reliable transit expands access to education and work.")

    def headings(document: DocumentObject) -> None:
        document.add_heading("Public Transit Case", level=1)
        document.add_heading("Accessibility", level=2)
        document.add_paragraph("Affordable transport reduces geographic exclusion.")

    def formatting(document: DocumentObject) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run("Bold claim.").bold = True
        paragraph.add_run(" Italic warrant.").italic = True
        paragraph.add_run(" Underlined impact.").underline = True
        highlighted = paragraph.add_run(" Highlighted weighing.")
        highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def table(document: DocumentObject) -> None:
        document.add_paragraph("Comparison follows.")
        table_object = document.add_table(rows=2, cols=2)
        table_object.cell(0, 0).text = "Policy"
        table_object.cell(0, 1).text = "Benefit"
        table_object.cell(1, 0).text = "Bus lanes"
        table_object.cell(1, 1).text = "Shorter travel times"
        document.add_paragraph("Comparison complete.")

    _save_document("Personal Debate Files/ordinary.docx", ordinary)
    _save_document("Personal Debate Files/formatting.docx", formatting)
    _save_document("Past Cases/2024/heading.docx", headings)
    _save_document("Other Debate Files/tables/table.docx", table)

    _write_text_pdf(
        FIXTURE_ROOT / "Other Debate Files/reference.PDF",
        "Extractable evidence says public transit improves access to opportunity.",
    )
    blank_pdf = PdfWriter()
    blank_pdf.add_blank_page(width=612, height=792)
    scan_path = FIXTURE_ROOT / "Scans/empty_scan.pdf"
    scan_path.parent.mkdir(parents=True, exist_ok=True)
    with scan_path.open("wb") as output:
        blank_pdf.write(output)

    markdown_path = FIXTURE_ROOT / "Non-Personal Files/brief.MD"
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(
        "# Synthetic Motion Brief\n\nPublic goods can justify collective investment.\n",
        encoding="utf-8",
    )
    text_path = FIXTURE_ROOT / "Nested/Research/notes.TXT"
    text_path.parent.mkdir(parents=True, exist_ok=True)
    text_path.write_text(
        "This is synthetic fixture text.\nIt contains no real debate archive material.\n",
        encoding="utf-8",
    )

    malformed_path = FIXTURE_ROOT / "Broken/malformed.docx"
    malformed_path.parent.mkdir(parents=True, exist_ok=True)
    malformed_path.write_text("This is intentionally not a zip archive.", encoding="utf-8")

    temporary_path = FIXTURE_ROOT / "Temp/~$temporary.docx"
    temporary_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path.write_text("Office lock-file fixture", encoding="utf-8")

    unsupported_path = FIXTURE_ROOT / "Other Debate Files/unsupported.csv"
    unsupported_path.write_text("kind,value\nfixture,ignored\n", encoding="utf-8")

    hidden_path = FIXTURE_ROOT / ".hidden/secret.txt"
    hidden_path.parent.mkdir(parents=True, exist_ok=True)
    hidden_path.write_text("Hidden fixture", encoding="utf-8")
    (FIXTURE_ROOT / ".DS_Store").write_bytes(b"fixture")

    def policy_structure(document: DocumentObject) -> None:
        document.add_heading("AD 1: Economic Mobility", level=1)
        document.add_heading("UQ:", level=2)
        document.add_paragraph("Transit access is currently unequal.")
        document.add_heading("L:", level=2)
        document.add_paragraph("Dedicated investment expands service.")
        document.add_heading("IL:", level=2)
        document.add_paragraph("Expanded service improves job access.")
        document.add_heading("IMPX:", level=2)
        document.add_paragraph("Economic mobility reduces entrenched poverty.")
        document.add_heading("DA 2: Inflation", level=1)
        document.add_heading("Uniqueness:", level=2)
        document.add_paragraph("Inflation is presently stable.")
        document.add_heading("Links:", level=2)
        document.add_paragraph("Rapid spending can increase demand.")
        document.add_heading("Internal Links:", level=2)
        document.add_paragraph("Demand can outpace productive capacity.")
        document.add_heading("Impacts:", level=2)
        document.add_paragraph("Price instability harms low-income households.")

    def harms_structure(document: DocumentObject) -> None:
        document.add_heading("Harms", level=1)
        document.add_paragraph("Current transport exclusion limits opportunity.")
        document.add_heading("Solvency", level=1)
        document.add_paragraph("Frequent service connects isolated communities.")
        document.add_heading("Impacts", level=1)
        document.add_paragraph("Mobility improves economic participation.")

    def contention_structure(document: DocumentObject) -> None:
        document.add_heading("Contention 1: Autonomy", level=1)
        document.add_heading("Claim:", level=2)
        document.add_paragraph("Accessible transport expands meaningful choice.")
        document.add_heading("Warrants:", level=2)
        document.add_paragraph("People need mobility to exercise available options.")
        document.add_heading("Impact:", level=2)
        document.add_paragraph("Autonomy is central to individual flourishing.")

    def value_structure(document: DocumentObject) -> None:
        document.add_heading("Value: Justice", level=1)
        document.add_paragraph("Justice evaluates whether institutions treat people fairly.")
        document.add_heading("Value Criterion: Equal Opportunity", level=1)
        document.add_paragraph("Prefer the world that equalizes access to social goods.")
        document.add_heading("Weighing Mechanism:", level=1)
        document.add_paragraph("Prioritize the least advantaged.")

    def fact_structure(document: DocumentObject) -> None:
        document.add_heading("Threshold of Truth", level=1)
        document.add_paragraph("The claim must be more likely true than false.")
        document.add_heading("Observation 1: Measurement", level=1)
        document.add_paragraph("Use independently replicated evidence.")
        document.add_heading("Claim", level=1)
        document.add_paragraph("The measured trend is persistent.")

    def theory_structure(document: DocumentObject) -> None:
        document.add_heading("Theory Shell", level=1)
        document.add_heading("Interpretation:", level=2)
        document.add_paragraph("Debaters must disclose advocacy before the round.")
        document.add_heading("Standards:", level=2)
        document.add_paragraph("Disclosure improves preparation and clash.")
        document.add_heading("Voters:", level=2)
        document.add_paragraph("Fairness and education justify the ballot.")

    def counter_theory_structure(document: DocumentObject) -> None:
        document.add_heading("Conditionality Theory", level=1)
        document.add_heading("Counter-Interp:", level=2)
        document.add_paragraph("One conditional advocacy is permissible.")
        document.add_heading("Counter-Standards:", level=2)
        document.add_paragraph("Limited flexibility improves testing.")
        document.add_heading("Competing Interpretations", level=2)
        document.add_paragraph("Prefer the interpretation with the best model.")
        document.add_heading("Reasonability", level=2)
        document.add_paragraph("Reject only practices that are clearly abusive.")

    def kritik_structure(document: DocumentObject) -> None:
        document.add_heading("Capitalism Kritik", level=1)
        document.add_heading("Framework:", level=2)
        document.add_paragraph("Evaluate the assumptions that organize advocacy.")
        document.add_heading("Link:", level=2)
        document.add_paragraph("The affirmative treats growth as an unquestioned good.")
        document.add_heading("Impact:", level=2)
        document.add_paragraph("Growth-first logic reproduces exploitation.")
        document.add_heading("Alt:", level=2)
        document.add_paragraph("Reject growth as the default measure of welfare.")

    def answers_structure(document: DocumentObject) -> None:
        document.add_heading("AT: Solvency", level=1)
        document.add_paragraph("Their mechanism cannot reach rural communities.")
        document.add_heading("A2 Inflation", level=1)
        document.add_paragraph("Targeted investment does not overheat aggregate demand.")
        document.add_heading("Answer To: Framework", level=1)
        document.add_paragraph("Consequences remain relevant to the ballot.")

    def weak_structure(document: DocumentObject) -> None:
        document.add_paragraph("Synthetic preface without an identified heading.")
        framing = document.add_paragraph()
        framing.add_run("ROUND FRAMING").bold = True
        document.add_paragraph("Compare structural access before marginal convenience.")
        highlighted = document.add_paragraph()
        run = highlighted.add_run("Long-Term Weighing")
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        document.add_paragraph("Durable access outweighs a temporary delay.")

    def nested_numbering(document: DocumentObject) -> None:
        document.add_heading("UQ", level=1)
        document.add_paragraph("1. Poverty")
        document.add_paragraph("Poverty remains geographically concentrated.")
        document.add_paragraph("a. Below the poverty line")
        document.add_paragraph("The synthetic baseline is illustrative only.")
        document.add_paragraph("i. Rural communities")
        document.add_paragraph("Distance compounds access barriers.")
        document.add_paragraph("2. Inflation")
        document.add_paragraph("Inflation remains within the target range.")

    def ordinary_prose(document: DocumentObject) -> None:
        document.add_paragraph(
            "The policy may have an impact on access, but this is an ordinary sentence."
        )
        document.add_paragraph(
            "The causal link depends on service frequency and should not become a heading."
        )
        document.add_paragraph(
            "Our framework for comparison considers both cost and geographic coverage."
        )
        document.add_paragraph(
            "The claim that every highlighted phrase is structural would be unreliable."
        )

    _save_structure_document("Case File Sandhu.docx", policy_structure)
    _save_structure_document("harms_solvency.docx", harms_structure)
    _save_structure_document("claim_warrant_impact.docx", contention_structure)
    _save_structure_document("value_case.docx", value_structure)
    _save_structure_document("fact_round.docx", fact_structure)
    _save_structure_document("Theory File - Sandhu.docx", theory_structure)
    _save_structure_document("counter_theory.docx", counter_theory_structure)
    _save_structure_document("capitalism_kritik.docx", kritik_structure)
    _save_structure_document("answers.docx", answers_structure)
    _save_structure_document("weak_headings.docx", weak_structure)
    _save_structure_document("nested_numbering.docx", nested_numbering)
    _save_structure_document("ordinary_prose.docx", ordinary_prose)


if __name__ == "__main__":
    generate()
