"""DOCX parsing with lightweight structure and run-format signals."""

from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from debate_engine.schemas import (
    DocumentFormat,
    ParsedBlock,
    ParsedDocument,
    ParseStatus,
)

_HEADING_STYLE = re.compile(r"^heading\s+(\d+)$", re.IGNORECASE)


def _heading_level(style_name: str | None) -> int | None:
    if not style_name:
        return None
    match = _HEADING_STYLE.fullmatch(style_name.strip())
    return int(match.group(1)) if match else None


def _format_signals(paragraphs: Iterable[Paragraph]) -> tuple[bool, bool, bool, bool]:
    runs = [run for paragraph in paragraphs for run in paragraph.runs]
    return (
        any(run.bold is True for run in runs),
        any(run.italic is True for run in runs),
        any(bool(run.underline) for run in runs),
        any(run.font.highlight_color is not None for run in runs),
    )


def _paragraph_block(paragraph: Paragraph, index: int) -> ParsedBlock:
    style_name = paragraph.style.name if paragraph.style is not None else None
    bold, italic, underline, highlight = _format_signals([paragraph])
    return ParsedBlock(
        text=paragraph.text,
        index=index,
        block_type="paragraph",
        style_name=style_name,
        heading_level=_heading_level(style_name),
        contains_bold=bold,
        contains_italic=italic,
        contains_underline=underline,
        contains_highlight=highlight,
    )


def _table_blocks(table: Table, table_index: int, start_index: int) -> list[ParsedBlock]:
    """Return table cells in row-major reading order.

    A cell is one block; internal paragraphs are separated by newlines. Merged
    cells can appear more than once through python-docx, so the same underlying
    XML cell is emitted only once per row.
    """
    blocks: list[ParsedBlock] = []
    next_index = start_index
    for row_index, row in enumerate(table.rows):
        seen_cells: set[int] = set()
        for column_index, cell in enumerate(row.cells):
            cell_identity = id(cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)

            paragraphs = list(cell.paragraphs)
            text = "\n".join(paragraph.text for paragraph in paragraphs)
            bold, italic, underline, highlight = _format_signals(paragraphs)
            style_names = {
                paragraph.style.name for paragraph in paragraphs if paragraph.style is not None
            }
            style_name = next(iter(style_names)) if len(style_names) == 1 else None
            blocks.append(
                ParsedBlock(
                    text=text,
                    index=next_index,
                    block_type="table_cell",
                    style_name=style_name,
                    heading_level=_heading_level(style_name),
                    contains_bold=bold,
                    contains_italic=italic,
                    contains_underline=underline,
                    contains_highlight=highlight,
                    table_index=table_index,
                    row_index=row_index,
                    column_index=column_index,
                )
            )
            next_index += 1
    return blocks


def _body_items(document: DocumentObject) -> Iterable[Paragraph | Table]:
    """Yield top-level paragraphs and tables in their XML document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


class DocxParser:
    """Extract DOCX body content without interpreting debate semantics."""

    extensions = frozenset({".docx"})

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(path)
        blocks: list[ParsedBlock] = []
        table_index = 0

        for item in _body_items(document):
            if isinstance(item, Paragraph):
                blocks.append(_paragraph_block(item, len(blocks)))
                continue
            blocks.extend(_table_blocks(item, table_index, len(blocks)))
            table_index += 1

        raw_text = "\n".join(block.text for block in blocks)
        warnings: list[str] = []
        if not raw_text.strip():
            warnings.append("Document contains no extractable body text.")

        return ParsedDocument(
            source_path=path,
            filename=path.name,
            detected_format=DocumentFormat.DOCX,
            raw_text=raw_text,
            blocks=blocks,
            warnings=warnings,
            parse_status=ParseStatus.PARTIAL if warnings else ParseStatus.SUCCESS,
        )
